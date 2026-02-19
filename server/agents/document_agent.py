"""Document Agent - read and analyze files (PDF, DOCX, TXT)."""

from __future__ import annotations

import asyncio
import os
import re
from typing import Any, Dict, List, Optional, Tuple

from .base_agent import BaseAgent, Task, AgentResponse, AgentStatus


class DocumentAgent(BaseAgent):
    def __init__(self, agent_id: str = "document_agent_v1"):
        super().__init__(
            agent_id=agent_id,
            name="Document Agent",
            description="Reads documents (PDF/DOCX/TXT), searches within them, and answers questions."
        )

    def can_handle_task(self, task: Task) -> bool:
        query = str(task.metadata.get("query", "")).lower()
        keywords = ["pdf", "document", "docx", "read file", "summarize file", "search file", "extract"]
        return any(k in query for k in keywords)

    async def process_task(self, task: Task) -> AgentResponse:
        self.status = AgentStatus.RUNNING
        task_id = task.metadata.get("parent_task_id") or task.id
        query = task.metadata.get("query", "")

        await self.emit_event("active_agent", task_id, {"agent": self.name})
        await self.emit_event("agent_state", task_id, {"state": "executing", "agent": self.name})
        await self.emit_event("agent_activity", task_id, {"agent": self.name, "message": "Processing document request"})

        file_path = self._extract_path(query)
        if not file_path:
            return AgentResponse(agent_id=self.agent_id, success=False, error="file_path_required")

        try:
            text, meta = await self._load_text(file_path)
            mode = self._infer_mode(query)
            if mode == "search":
                matches = self._search_text(text, query)
                return AgentResponse(agent_id=self.agent_id, success=True, result={
                    "action": "search",
                    "file": file_path,
                    "matches": matches,
                    "meta": meta,
                })
            if mode == "extract":
                return AgentResponse(agent_id=self.agent_id, success=True, result={
                    "action": "extract",
                    "file": file_path,
                    "content": text[:8000],
                    "meta": meta,
                })

            llm = self.get_llm(task)
            prompt = f"Answer the user's request using the document content.\nUser request: {query}\nDocument content:\n{text[:12000]}"
            response = await llm.ainvoke(prompt)
            answer = response.content if hasattr(response, "content") else str(response)
            return AgentResponse(agent_id=self.agent_id, success=True, result={
                "action": mode,
                "file": file_path,
                "answer": answer,
                "meta": meta,
            })
        except Exception as exc:
            return AgentResponse(agent_id=self.agent_id, success=False, error=str(exc))
        finally:
            self.status = AgentStatus.COMPLETED

    def _extract_path(self, text: str) -> str:
        windows_match = re.search(r"([A-Za-z]:\\[^\"'\s]+)", text)
        if windows_match:
            return windows_match.group(1)
        unix_match = re.search(r"(/[^\"'\s]+)", text)
        if unix_match:
            return unix_match.group(1)
        return ""

    def _infer_mode(self, query: str) -> str:
        q = query.lower()
        if "search" in q or "find" in q:
            return "search"
        if "extract" in q or "show" in q or "dump" in q:
            return "extract"
        if "summarize" in q or "summary" in q:
            return "summarize"
        return "qa"

    async def _load_text(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        if not os.path.exists(file_path):
            raise RuntimeError("file_not_found")
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return await asyncio.to_thread(self._read_pdf, file_path)
        if ext in {".docx", ".doc"}:
            return await asyncio.to_thread(self._read_docx, file_path)
        return await asyncio.to_thread(self._read_text, file_path)

    def _read_pdf(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n".join(pages)
        return text, {"pages": len(reader.pages), "type": "pdf"}

    def _read_docx(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        from docx import Document
        doc = Document(file_path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text)
        return text, {"paragraphs": len(doc.paragraphs), "type": "docx"}

    def _read_text(self, file_path: str) -> Tuple[str, Dict[str, Any]]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return text, {"type": "text"}

    def _search_text(self, text: str, query: str) -> List[Dict[str, Any]]:
        terms = [t for t in re.findall(r"\w+", query.lower()) if len(t) > 2]
        matches = []
        for term in terms[:5]:
            for m in re.finditer(term, text.lower()):
                start = max(0, m.start() - 50)
                end = min(len(text), m.end() + 50)
                matches.append({"term": term, "snippet": text[start:end]})
                if len(matches) >= 20:
                    return matches
        return matches
